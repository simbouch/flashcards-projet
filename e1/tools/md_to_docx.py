from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RED_TODO = RGBColor(0xC0, 0x00, 0x00)


def _add_field(paragraph, instr: str) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    paragraph._p.append(fld)


def _add_complex_field(paragraph, instr: str, *, placeholder: str = "") -> None:
    """Add a Word 'complex' field.

    This is more compatible for fields like TOC that otherwise look empty until updated.
    """

    r_begin = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r_begin._r.append(fld_begin)

    r_instr = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    r_instr._r.append(instr_el)

    r_sep = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    r_sep._r.append(fld_sep)

    if placeholder:
        paragraph.add_run(placeholder)

    r_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_end._r.append(fld_end)


def _enable_update_fields_on_open(doc: Document) -> None:
    """Ask Word to update fields (TOC, PAGE, etc.) when the document is opened."""

    settings = doc.settings.element
    if settings.xpath("./w:updateFields"):
        return
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def _tokenize_inline(text: str):
    pat = re.compile(r"(`[^`]+`|\*\*.+?\*\*|\*.+?\*)")
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            yield ("text", text[pos : m.start()])
        tok = m.group(0)
        if tok.startswith("`"):
            yield ("code", tok[1:-1])
        elif tok.startswith("**"):
            yield ("bold", tok[2:-2])
        else:
            yield ("italic", tok[1:-1])
        pos = m.end()
    if pos < len(text):
        yield ("text", text[pos:])


def _add_runs(p, text: str, *, todo: bool = False) -> None:
    for kind, chunk in _tokenize_inline(text):
        if not chunk:
            continue
        run = p.add_run(chunk)
        if todo:
            run.font.color.rgb = RED_TODO
            run.bold = True
        if kind == "code":
            run.font.name = "Consolas"
        elif kind == "bold":
            run.bold = True
        elif kind == "italic":
            run.italic = True


def _add_inline_with_todos(p, text: str) -> None:
    parts = re.split(r"(\[\[.*?\]\])", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("[[") and part.endswith("]]"):
            _add_runs(p, part[2:-2], todo=True)
        else:
            _add_runs(p, part)


def build_doc(md_text: str) -> Document:
    doc = Document()

    # Page setup (trainer template: A4, standard margins)
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    # Ensure fields (TOC/page numbers) are updated on open.
    _enable_update_fields_on_open(doc)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Footer: page X / Y
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.alignment = 1
    fp.add_run("Page ")
    _add_field(fp, "PAGE")
    fp.add_run(" / ")
    _add_field(fp, "NUMPAGES")

    def _split_md_table_row(row_line: str) -> list[str]:
        parts = [c.strip() for c in row_line.strip().strip("|").split("|")]
        return [p for p in parts if p != ""]

    def _is_md_table_sep(row_line: str) -> bool:
        s = row_line.strip().strip("|").replace(" ", "")
        # e.g. "---|---|---" or ":---|---:|---"
        return bool(s) and all(set(col) <= set("-:") and "-" in col for col in s.split("|"))


    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    first_hr = True
    toc_inserted = False
    last_blank = False

    while i < len(lines):
        line = lines[i].rstrip("\r")

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                for j, cl in enumerate(code_lines):
                    r = p.add_run(cl)
                    r.font.name = "Consolas"
                    r.font.size = Pt(10)
                    if j < len(code_lines) - 1:
                        r.add_break()
                in_code = False
            last_blank = False
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Markdown tables (pipe tables)
        if "|" in line and i + 1 < len(lines) and _is_md_table_sep(lines[i + 1]):
            header = _split_md_table_row(line)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines):
                row_line = lines[i].rstrip("\r")
                if not row_line.strip() or "|" not in row_line:
                    break
                rows.append(_split_md_table_row(row_line))
                i += 1

            cols = max(1, len(header))
            tbl = doc.add_table(rows=1 + len(rows), cols=cols)
            try:
                tbl.style = "Table Grid"
            except Exception:
                pass

            # header row
            for c, txt in enumerate(header[:cols]):
                p = tbl.cell(0, c).paragraphs[0]
                p.text = ""
                _add_inline_with_todos(p, txt)
                for run in p.runs:
                    run.bold = True

            # body
            for r_idx, row in enumerate(rows, start=1):
                for c in range(cols):
                    p = tbl.cell(r_idx, c).paragraphs[0]
                    p.text = ""
                    _add_inline_with_todos(p, row[c] if c < len(row) else "")

            last_blank = False
            continue


        # Explicit TOC marker (allows Remerciements before TOC, as in trainer template)
        if line.strip() == "[[TOC]]" and not toc_inserted:
            doc.add_page_break()

            # Use a Title-style paragraph (not a Heading) so it doesn't appear inside the TOC itself.
            p_title = doc.add_paragraph("Sommaire")
            try:
                p_title.style = "Title"
            except Exception:
                pass

            p_toc = doc.add_paragraph()
            _add_complex_field(
                p_toc,
                'TOC \\o "1-4" \\h \\u',
                placeholder="(Mettre à jour le champ pour afficher le sommaire)",
            )
            doc.add_page_break()
            toc_inserted = True
            last_blank = False
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
            last_blank = False
            i += 1
            continue

        if line.strip() == "---":
            if first_hr:

                # End of cover
                doc.add_page_break()
                first_hr = False
            else:
                doc.add_paragraph()
            last_blank = True
            i += 1
            continue

        if not line.strip():
            if not last_blank:
                doc.add_paragraph()
            last_blank = True
            i += 1
            continue

        stripped = line.lstrip(" ")
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_with_todos(p, stripped[2:].strip())
        else:
            p = doc.add_paragraph()
            _add_inline_with_todos(p, line.strip())
        last_blank = False

        i += 1

    return doc


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="inp", default="e1/docs/rapport_e1.md")
    ap.add_argument("--out", dest="out", default="e1/docs/rapport_e1_KHRIBECH_B.docx")
    args = ap.parse_args()

    inp = Path(args.inp)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = build_doc(inp.read_text(encoding="utf-8"))
    doc.save(out)
    print(f"[OK] Wrote: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

